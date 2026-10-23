#!/usr/bin/env python3
"""Export the effective runtime bank as lesson-wise Word moderation packs.

Must be run against the assembled /app application after ALL V5/V6 patches.
Never fabricate missing answer explanations or expert verification statuses.
"""
import argparse
import json
import re
import hashlib
from collections import Counter, OrderedDict
from datetime import datetime, timezone
from pathlib import Path
from zipfile import ZIP_DEFLATED, ZipFile

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.shared import Cm, Pt, RGBColor

CHAPTER_RANK = {"Chương I": 1, "Chương II": 2, "Chương III": 3, "Chương IV": 4}
TYPE_LIST = (("mcq", "PHẦN I. TRẮC NGHIỆM NHIỀU LỰA CHỌN"),
             ("tf", "PHẦN II. TRẮC NGHIỆM ĐÚNG / SAI"),
             ("short", "PHẦN III. TRẢ LỜI NGẮN"))
LEVELS = {"NB": "Nhận biết", "TH": "Thông hiểu",
          "VD": "Vận dụng", "VDC": "Vận dụng cao"}
LETTERS = "ABCD"


def text(value):
    """Make stored math text editable in Word; never silently recalculate it."""
    s = str(value if value is not None else "")
    # Remove display-markup wrappers, not the mathematical contents.
    for a, b in ((r"\(", ""), (r"\)", ""), (r"\[", ""), (r"\]", ""),
                 (r"\times", "×"), (r"\cdot", "·"), (r"\Delta", "Δ"),
                 (r"\alpha", "α"), (r"\beta", "β"), (r"\gamma", "γ"),
                 (r"\lambda", "λ"), (r"\pi", "π"), (r"\theta", "θ"),
                 (r"\Omega", "Ω"), (r"\mu", "μ"), (r"\infty", "∞"),
                 (r"\mathrm", ""), (r"\text", "")):
        s = s.replace(a, b)
    return s.strip()


def lesson_number(lesson):
    m = re.search(r"\d+", str(lesson))
    return int(m.group()) if m else 999


def lesson_title(data, chapter, lesson):
    record = data.get("chapters", {}).get(chapter, {}).get("lessons", {})
    obj = record.get(lesson) if isinstance(record, dict) else None
    if isinstance(obj, str):
        return obj.strip()
    if isinstance(obj, dict):
        for k in ("name", "title", "topic", "label"):
            if obj.get(k):
                return str(obj[k]).strip()
    return ""


def add_text(doc, body, *, bold_prefix=None):
    p = doc.add_paragraph(style="Normal")
    if bold_prefix:
        p.add_run(bold_prefix).bold = True
    p.add_run(text(body))
    return p


def configure(doc, title):
    section = doc.sections[0]
    section.top_margin = Cm(1.75)
    section.bottom_margin = Cm(1.55)
    section.left_margin = Cm(2.0)
    section.right_margin = Cm(1.8)
    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(4)
    normal.paragraph_format.line_spacing = 1.14
    for h, size in (("Title", 15), ("Heading 1", 12.5), ("Heading 2", 11.5)):
        st = doc.styles[h]
        st.font.name = "Times New Roman"
        st.font.size = Pt(size)
        st.font.bold = True
        st.font.color.rgb = RGBColor(30, 54, 82)
    hdr = section.header.paragraphs[0]
    hdr.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    hdr.add_run("PHY-AI DIAGNOSE  |  BẢN KIỂM DUYỆT NỘI BỘ").italic = True
    hdr.runs[0].font.size = Pt(8)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.add_run("Bản phục vụ giáo viên kiểm duyệt — chưa xác nhận chất lượng thực nghiệm").italic = True
    footer.runs[0].font.size = Pt(8)
    doc.add_heading(title, 0)


def question_head(doc, q, serial):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.keep_with_next = True
    r = p.add_run(f"Câu {serial:02d}.  ")
    r.bold = True
    meta = (f"[{q['id']}]  •  {LEVELS.get(q.get('level'), q.get('level','Chưa rõ'))}"
            f"  •  Lỗi: {q.get('error') or '—'}"
            f"  •  {'Chẩn đoán' if q.get('diagnostic') is True else 'Ôn tập'}")
    m = p.add_run(meta)
    m.italic = True
    m.font.size = Pt(9)
    return p


def write_question(doc, q, serial):
    question_head(doc, q, serial)
    add_text(doc, q["prompt"])
    if q["type"] == "mcq":
        for i, option in enumerate(q["options"]):
            add_text(doc, option, bold_prefix=LETTERS[i] + ". ")
    elif q["type"] == "tf":
        for i, st in enumerate(q["statements"]):
            add_text(doc, st["text"], bold_prefix="abcd"[i] + ") ")
    else:
        add_text(doc, "Trả lời: ..................................................................")


def display_number(x):
    if isinstance(x, float):
        return format(x, ".12g").replace(".", ",")
    return str(x).replace(".", ",") if isinstance(x, int) else text(x)


def write_answer(doc, q, serial):
    question_head(doc, q, serial)
    if q["type"] == "mcq":
        index = q["answer"]
        add_text(doc, f"{LETTERS[index]}. {q['options'][index]}",
                 bold_prefix="Đáp án: ")
    elif q["type"] == "tf":
        key = " – ".join(f"{'abcd'[i]}) {'Đúng' if st['answer'] else 'Sai'}"
                         for i, st in enumerate(q["statements"]))
        add_text(doc, key, bold_prefix="Đáp án: ")
        for i, st in enumerate(q["statements"]):
            txt = f"{'Đúng' if st['answer'] else 'Sai'} — {st['text']}"
            if st.get("explain"):
                txt += " | " + st["explain"]
            add_text(doc, txt, bold_prefix=f"{'abcd'[i]}) ")
    else:
        val = display_number(q["answer"])
        unit = text(q.get("unit", ""))
        add_text(doc, f"{val} {unit}".strip(), bold_prefix="Đáp án: ")
        add_text(doc, display_number(q.get("tolerance", "")),
                 bold_prefix="Sai số chấp nhận của hệ thống: ")
    add_text(doc, q.get("explain") or "Chưa có giải thích riêng trong dữ liệu.",
             bold_prefix="Giải thích đang lưu: ")
    if q.get("source_page"):
        add_text(doc, str(q["source_page"]), bold_prefix="Trang SGK được gắn nguồn: ")
    if q.get("verified") is False:
        add_text(doc, "CHƯA được giáo viên thẩm định độc lập", bold_prefix="CẢNH BÁO: ")
    doc.add_paragraph(
        "Đánh giá của GV: ☐ Đạt  ☐ Sửa kiến thức  ☐ Sửa đáp án  ☐ Sửa diễn đạt  ☐ Loại"
    )
    doc.add_paragraph("Nhận xét / phương án sửa: ........................................................................................")


def document_for(data, grouped, chapter, lesson, answer=False):
    items = grouped[(chapter, lesson)]
    doc = Document()
    task_name = "ĐÁP ÁN – GIẢI THÍCH – PHIẾU KIỂM DUYỆT" if answer else "NGÂN HÀNG CÂU HỎI"
    configure(doc, task_name)
    title = lesson_title(data, chapter, lesson)
    doc.add_heading(f"{chapter}  |  {lesson}" + (f": {title}" if title else ""), 1)
    stat = Counter(q["type"] for q in items)
    doc.add_paragraph(
        f"Tổng {len(items)} câu  |  Nhiều lựa chọn {stat['mcq']}  |  "
        f"Đúng/Sai {stat['tf']}  |  Trả lời ngắn {stat['short']}"
    )
    doc.add_paragraph("Phiên bản ngân hàng: " + data["version"])
    doc.add_paragraph("Mã câu và thứ tự đáp án giữ nguyên bản hiệu lực; không hoán vị khi xuất.")
    if answer:
        doc.add_paragraph(
            "Lưu ý: 'Đáp án' và 'Giải thích' là nội dung đang lưu trong hệ thống; "
            "chưa mặc nhiên được giáo viên chứng nhận đúng."
        )
    serial = 0
    for typ, heading in TYPE_LIST:
        in_type = [q for q in items if q["type"] == typ]
        if not in_type:
            continue
        doc.add_heading(f"{heading} ({len(in_type)} câu)", 2)
        for q in in_type:
            serial += 1
            (write_answer if answer else write_question)(doc, q, serial)
    return doc


def write_index(data, grouped, out_dir, issue_list):
    doc = Document()
    configure(doc, "DANH MỤC VÀ PHIẾU BÀN GIAO KIỂM DUYỆT 1.000 CÂU")
    doc.add_paragraph("Phiên bản: " + data["version"])
    doc.add_paragraph(
        "Bộ hồ sơ gồm tệp câu hỏi (không lộ đáp án) và tệp đáp án/giải thích "
        "cho mỗi bài. Số thứ tự được đánh lại trong từng bài; mã ID là khóa đối chiếu."
    )
    rows = doc.add_table(rows=1, cols=6)
    rows.style = "Table Grid"
    for cell, value in zip(rows.rows[0].cells,
                           ("Chương", "Bài", "Tổng", "MCQ", "Đ/S", "Ngắn")):
        cell.text = value
    for chapter, lesson in grouped:
        items = grouped[(chapter, lesson)]
        c = Counter(q["type"] for q in items)
        cells = rows.add_row().cells
        for cell, val in zip(cells, (chapter, lesson, len(items), c['mcq'], c['tf'], c['short'])):
            cell.text = str(val)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    doc.add_heading("Trình tự giáo viên kiểm duyệt", 1)
    for label in (
        "Đọc câu hỏi độc lập, đối chiếu nội dung bài học trong SGK.",
        "Kiểm tra tính khoa học, điều kiện áp dụng, đơn vị, dấu và làm tròn.",
        "Đối chiếu đáp án và toàn bộ phương án nhiễu / mệnh đề Đúng–Sai.",
        "Xác minh mã lỗi tư duy, độ khó, tính độc lập của các câu chẩn đoán.",
        "Ghi sửa đổi theo mã câu, không chỉ theo số thứ tự.",
        "Xác nhận độc lập các câu có trạng thái verified=false."
    ):
        doc.add_paragraph(label, style="List Bullet")
    doc.add_heading("Những vấn đề cần chú ý", 1)
    doc.add_paragraph(
        "Hồ sơ này chỉ sao xuất dữ liệu đang có. Không có việc tự động chữa "
        "những câu sai nội dung; các kết luận khoa học cuối cùng do giáo viên phê duyệt."
    )
    doc.add_paragraph(f"Số câu gắn verified=false: {len(issue_list)}.")
    if issue_list:
        doc.add_paragraph("Các ID chờ xác minh: " + ", ".join(issue_list))
    path = out_dir / "00_DANH_MUC_VA_HUONG_DAN_KIEM_DUYET.docx"
    doc.save(path)
    return path


def export(data, out_dir):
    questions = data["questions"]
    if len(questions) != 1000:
        raise ValueError(f"Require 1000 questions, got {len(questions)}")
    if len({q["id"] for q in questions}) != 1000:
        raise ValueError("Duplicate question IDs; refusing to export")
    counts = Counter(q["type"] for q in questions)
    if counts != {"mcq":532, "tf":218, "short":250}:
        raise ValueError(f"Unexpected type distribution: {counts}")
    chapters = Counter(q["chapter"] for q in questions)
    if chapters != {"Chương I":250, "Chương II":250,"Chương III":250,"Chương IV":250}:
        raise ValueError(f"Unexpected chapter distribution: {chapters}")
    grouped = {}
    for q in questions:
        key = (q["chapter"], q["lesson"])
        grouped.setdefault(key, []).append(q)
    grouped = OrderedDict(sorted(grouped.items(),key=lambda it:(
        CHAPTER_RANK.get(it[0][0], 999), lesson_number(it[0][1]))))
    out_dir.mkdir(parents=True, exist_ok=True)
    files = []
    manifest = []
    issues = [q["id"] for q in questions if q.get("verified") is False]
    for chapter, lesson in grouped:
        prefix = f"Chuong_{CHAPTER_RANK[chapter]:02d}_Bai_{lesson_number(lesson):02d}"
        cdir = out_dir / f"Chuong_{CHAPTER_RANK[chapter]:02d}"
        cdir.mkdir(exist_ok=True)
        for answer, suffix in ((False, "CAU_HOI"), (True, "DAP_AN_KIEM_DUYET")):
            doc = document_for(data, grouped, chapter, lesson, answer=answer)
            path = cdir / f"{prefix}_{suffix}.docx"
            doc.save(path)
            files.append(path)
        t = Counter(q["type"] for q in grouped[(chapter,lesson)])
        manifest.append({"chapter":chapter, "lesson":lesson, "count":len(grouped[(chapter,lesson)]),
                         "types":dict(t), "question_ids":[q["id"] for q in grouped[(chapter,lesson)]]})
    files.append(write_index(data, grouped, out_dir, issues))
    report = {
        "bank_version":data["version"], "question_count":len(questions),
        "distinct_ids":len({q["id"] for q in questions}),
        "type_counts":dict(counts), "chapter_counts":dict(chapters),
        "lesson_count":len(grouped), "docx_count":len(files),
        "unverified_question_ids":issues, "lessons":manifest
    }
    report_path=out_dir/"00_BAO_CAO_DOI_SOAT.json"
    report_path.write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding="utf-8")
    archive=out_dir.parent/"PHY_AI_DIAGNOSE_1000_CAU_THEO_BAI_KIEM_DUYET.zip"
    with ZipFile(archive,"w",compression=ZIP_DEFLATED,compresslevel=6) as z:
        for file in files+[report_path]:
            z.write(file,arcname=str(file.relative_to(out_dir.parent)))
    with ZipFile(archive) as z:
        docx_names=[n for n in z.namelist() if n.endswith(".docx")]
        if len(docx_names)!=2*len(grouped)+1:
            raise ValueError("Archive does not contain each question and answer Word file")
        for filename in docx_names:
            data_bytes=z.read(filename)
            if not data_bytes.startswith(b"PK"):
                raise ValueError("Corrupt Word member: "+filename)
    print(json.dumps({
        "version":data["version"], "questions":len(questions),
        "lessons":len(grouped), "word_files":len(files),
        "unverified_questions":len(issues),
        "archive":str(archive), "archive_bytes":archive.stat().st_size,
        "sha256":hashlib.sha256(archive.read_bytes()).hexdigest()
    },ensure_ascii=False,indent=2))


if __name__ == "__main__":
    p=argparse.ArgumentParser()
    p.add_argument("--out",type=Path,required=True)
    p.add_argument("--app-root",type=Path,default=Path("/app"))
    args=p.parse_args()
    import sys
    sys.path.insert(0,str(args.app_root))
    from app.question_bank import bank
    export(bank.data,args.out)
