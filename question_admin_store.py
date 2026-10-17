"""Versioned PostgreSQL question-management store; immutable built-in question bank is untouched."""
from __future__ import annotations
import io
import json
import re
from datetime import datetime, timezone
from uuid import uuid4
from sqlalchemy import Column, String, Integer, DateTime, JSON
from .database import Base

MAX_IMPORT_BYTES=3*1024*1024
HEADERS=("source_id","chapter","lesson","type","prompt","context","options","statements","answer","explain","error","diagnostic","tolerance","unit")

def utcnow():
    return datetime.now(timezone.utc)

class ManagedQuestion(Base):
    __tablename__="managed_question_v1"
    id=Column(String(40),primary_key=True)
    source_id=Column(String(80),nullable=True,index=True)
    draft_json=Column(JSON,nullable=False,default=dict)
    published_json=Column(JSON,nullable=True)
    history_json=Column(JSON,nullable=False,default=list)
    status=Column(String(32),nullable=False,default="draft")
    revision=Column(Integer,nullable=False,default=1)
    published_revision=Column(Integer,nullable=False,default=0)
    editor_id=Column(String(80),nullable=False)
    approver_id=Column(String(80),nullable=True)
    created_at=Column(DateTime(timezone=True),default=utcnow,nullable=False)
    updated_at=Column(DateTime(timezone=True),default=utcnow,onupdate=utcnow,nullable=False)
    published_at=Column(DateTime(timezone=True),nullable=True)

def new_id():
    return "QMG-"+uuid4().hex[:16].upper()

def version_id(id_,revision):
    return f"{id_}-v{revision}"

def _required_string(v,name,limit):
    if not isinstance(v,str) or not v.strip() or len(v.strip())>limit:
        raise ValueError(f"{name}: yêu cầu văn bản không rỗng, tối đa {limit} ký tự.")
    stripped=v.strip()
    if re.search(r"<\\s*(?:/?[A-Za-z][A-Za-z0-9]*|!--)", stripped):
        raise ValueError(f"{name}: không cho phép mã HTML trong nội dung câu hỏi.")
    return stripped

def normalize_question(payload,bank):
    """Reject unknown fields and any unreviewed ontology/scope mapping."""
    if not isinstance(payload,dict):
        raise ValueError("Câu hỏi phải là một đối tượng JSON.")
    allowed={"chapter","lesson","type","prompt","context","options","statements","answer","explain","error","diagnostic","tolerance","unit"}
    extras=set(payload)-allowed
    if extras: raise ValueError("Trường không hỗ trợ: "+", ".join(sorted(extras)))
    q={}
    q["chapter"]=_required_string(payload.get("chapter"),"Chương",40)
    chapters=bank.data.get("chapters") or {}
    if q["chapter"] not in chapters: raise ValueError("Chương không tồn tại trong SGK Vật lí 12.")
    q["lesson"]=_required_string(payload.get("lesson"),"Bài học",48)
    if q["lesson"] not in chapters[q["chapter"]].get("lessons",[]):
        raise ValueError("Bài học không thuộc chương đã chọn.")
    typ=_required_string(payload.get("type"),"Loại câu",10)
    if typ not in ("mcq","tf","short"): raise ValueError("Loại câu phải là mcq, tf hoặc short.")
    q["type"]=typ
    q["prompt"]=_required_string(payload.get("prompt"),"Câu dẫn",2500)
    q["explain"]=_required_string(payload.get("explain"),"Giải thích đáp án",3500)
    if payload.get("context"):
        q["context"]=_required_string(payload["context"],"Ngữ cảnh",1500)
    diagnostic=payload.get("diagnostic",False)
    if type(diagnostic) is not bool: raise ValueError("diagnostic phải là true hoặc false.")
    q["diagnostic"]=diagnostic
    err=payload.get("error")
    if err:
        err=_required_string(err,"Mã lỗi",60).upper()
        meta=bank.error_by_id.get(err)
        if not meta: raise ValueError("Mã lỗi tư duy chưa tồn tại trong ontology.")
        if meta.get("chapter") not in (None,q["chapter"]) or meta.get("lesson") not in (None,q["lesson"]):
            raise ValueError("Mã lỗi tư duy không thuộc chương/bài đã chọn.")
        q["error"]=err
    elif diagnostic:
        raise ValueError("Câu chẩn đoán bắt buộc có mã lỗi tư duy đã kiểm chứng.")
    if typ=="mcq":
        options=payload.get("options")
        if not isinstance(options,list) or len(options)!=4:
            raise ValueError("Trắc nghiệm phải có đúng 4 phương án.")
        options=[_required_string(x,"Phương án",700) for x in options]
        if len(set(options))!=4: raise ValueError("Các phương án phải khác nhau.")
        q["options"]=options
        answer=payload.get("answer")
        if isinstance(answer,str) and answer.strip().upper() in ("A","B","C","D"):
            answer="ABCD".index(answer.strip().upper())
        if type(answer) is not int or answer not in range(4):
            raise ValueError("Đáp án MCQ phải là A–D hoặc chỉ số 0–3.")
        q["answer"]=answer
    elif typ=="tf":
        statements=payload.get("statements")
        if not isinstance(statements,list) or len(statements)!=4:
            raise ValueError("Đúng/Sai cần 4 mệnh đề có đáp án và lời giải.")
        output=[]
        for item in statements:
            if not isinstance(item,dict) or type(item.get("answer")) is not bool:
                raise ValueError("Mệnh đề Đúng/Sai phải có answer true/false.")
            st={"text":_required_string(item.get("text"),"Mệnh đề",750),
                "answer":item["answer"],
                "explain":_required_string(item.get("explain"),"Giải thích mệnh đề",1000)}
            output.append(st)
        q["statements"]=output
    else:
        try:
            if type(payload.get("answer")) not in (int,float,str) or type(payload.get("answer")) is bool:
                raise ValueError("Đáp án phải là số.")
            ans=float(str(payload.get("answer")).replace(",","."))
            tol=float(str(payload.get("tolerance")).replace(",","."))
        except (TypeError,ValueError,OverflowError) as exc:
            raise ValueError("Trả lời ngắn yêu cầu đáp án và dung sai dạng số.") from exc
        import math
        if not math.isfinite(ans) or not math.isfinite(tol) or tol<=0 or tol>1e8:
            raise ValueError("Đáp án/dung sai không hợp lệ.")
        q["answer"]=ans;q["tolerance"]=tol
        if payload.get("unit"):
            q["unit"]=_required_string(payload["unit"],"Đơn vị",40)
    return q

def read_excel(content,bank):
    """Strict XLSX import: no formulas, macros, hidden rows, or partial writes."""
    if not isinstance(content,bytes) or len(content)>MAX_IMPORT_BYTES or len(content)<20:
        raise ValueError("Chỉ nhận .xlsx hợp lệ tối đa 3 MB.")
    import openpyxl
    try:
        book=openpyxl.load_workbook(io.BytesIO(content),read_only=False,data_only=False,keep_links=False)
    except Exception as exc:
        raise ValueError("Không đọc được tệp XLSX hợp lệ.") from exc
    if len(book.worksheets)!=1:
        raise ValueError("Excel chỉ được có một sheet Câu hỏi.")
    ws=book.active
    if ws.max_row>301 or ws.max_column>len(HEADERS):
        raise ValueError("Mỗi lần tối đa 300 câu, đúng 14 cột theo biểu mẫu.")
    actual=[str(ws.cell(1,i).value or "").strip() for i in range(1,len(HEADERS)+1)]
    if tuple(actual)!=HEADERS:
        raise ValueError("Tên hoặc thứ tự 14 cột không đúng biểu mẫu.")
    out=[]
    for row_i in range(2,ws.max_row+1):
        if ws.row_dimensions[row_i].hidden:
            raise ValueError(f"Dòng {row_i} đang ẩn; hãy kiểm tra trước khi nhập.")
        cells=[ws.cell(row_i,i) for i in range(1,len(HEADERS)+1)]
        if all(c.value is None for c in cells): continue
        if any(c.data_type=="f" for c in cells):
            raise ValueError(f"Dòng {row_i}: không cho phép công thức Excel.")
        row={key:cell.value for key,cell in zip(HEADERS,cells)}
        try:
            value={}
            for k in ("chapter","lesson","type","prompt","context","explain","error","unit"):
                if row[k] is not None:value[k]=str(row[k])
            source_id=str(row["source_id"]).strip() if row["source_id"] is not None else None
            typ=value.get("type")
            if typ=="mcq":
                value["options"]=str(row["options"] or "").split("|")
                value["answer"]=row["answer"]
            elif typ=="tf":
                value["statements"]=json.loads(row["statements"] or "[]")
            else:
                value["answer"]=row["answer"];value["tolerance"]=row["tolerance"]
            dv=row["diagnostic"]
            if dv is None:value["diagnostic"]=False
            elif type(dv) is bool:value["diagnostic"]=dv
            elif str(dv).strip().lower() in ("true","1","có","yes"):value["diagnostic"]=True
            elif str(dv).strip().lower() in ("false","0","không","no"):value["diagnostic"]=False
            else:raise ValueError("Cột diagnostic chỉ nhận true/false.")
            out.append((source_id,normalize_question(value,bank)))
        except (ValueError,TypeError,json.JSONDecodeError) as exc:
            raise ValueError(f"Dòng {row_i}: {exc}") from exc
    if not out:raise ValueError("Excel không có câu hỏi.")
    originals=[src for src,_ in out if src]
    if len(originals)!=len(set(originals)):
        raise ValueError("Một file không được sửa cùng một mã câu gốc nhiều lần.")
    return out

def workbook_template(bank):
    import openpyxl
    from openpyxl.styles import PatternFill,Font,Alignment
    w=openpyxl.Workbook();s=w.active;s.title="Cau hoi"
    s.append(list(HEADERS))
    for j in range(1,len(HEADERS)+1):
        c=s.cell(1,j);c.fill=PatternFill("solid",fgColor="1E4E79")
        c.font=Font(color="FFFFFF",bold=True)
        c.alignment=Alignment(wrap_text=True,vertical="center")
        s.column_dimensions[c.column_letter].width=34 if j in (5,6,7,8,10) else 19
    s.freeze_panes="D2";s.auto_filter.ref="A1:N301";s.row_dimensions[1].height=31
    a=io.BytesIO();w.save(a);return a.getvalue()
